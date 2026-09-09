#!/usr/bin/env python3
"""
Apply Word heading styles from a sample outline onto unformatted .docx files.

The Word-to-Elementor plugin reads Heading 1 / 2 / 3 (not bold-only text).
Sample Format.docx is the outline this script copies:

  Heading 1  page title
  Normal     intro paragraphs
  Heading 2  section title (Services, Why Choose Us, Process, FAQ, Closing)
  Heading 3  card / step / FAQ question
  Normal     card body / FAQ answer / closing copy
"""

from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

HEADING_STYLE = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3"}
BODY_STYLE = "Normal"

DEFAULT_SECTION_TITLES = {
    "services": "Our Services",
    "why": "Why Choose Us",
    "process": "Our Process",
    "faq": "FAQs",
    "closing": "Closing Section",
}

SECTION_ORDER = ("services", "why", "process", "faq", "closing")
INPUT_DIR = Path("unformatted")
OUTPUT_DIR = Path("formatted")
SAMPLE_DIR = Path("sample formats")
DEFAULT_SAMPLE_NAME = "Sample Format.docx"


@dataclass
class Item:
    title: str
    bodies: list[str] = field(default_factory=list)


@dataclass
class Outline:
    title: str = ""
    intro: list[str] = field(default_factory=list)
    section_titles: dict[str, str] = field(default_factory=dict)
    items: dict[str, list[Item]] = field(default_factory=dict)
    closing_bodies: list[str] = field(default_factory=list)

    def __post_init__(self) -> None:
        for key in SECTION_ORDER:
            self.section_titles.setdefault(key, "")
            self.items.setdefault(key, [])


def paragraph_text(paragraph: Paragraph) -> str:
    return "".join(run.text or "" for run in paragraph.runs).strip()


def heading_level(paragraph: Paragraph) -> int:
    style = paragraph.style
    name = (style.name if style is not None else "") or ""
    match = re.search(r"heading\s*([1-3])", name, re.I)
    if match:
        return int(match.group(1))
    style_id = getattr(style, "style_id", "") or ""
    match = re.search(r"heading\s*([1-3])", style_id, re.I)
    if match:
        return int(match.group(1))
    if style_id in {"1", "2", "3"}:
        return int(style_id)
    return 0


def normalize(text: str) -> str:
    text = text.lower()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def classify_section_title(text: str) -> str | None:
    if text.strip().endswith("?"):
        return None
    key = normalize(text)
    if not key:
        return None
    words = key.split()
    if key in {"faq", "faqs", "frequently asked questions"} or key.endswith(" faq"):
        return "faq"
    if key in {"process", "our process"} or key.endswith(" process"):
        return "process"
    if key.startswith("why choose") or key.startswith("why do people") or key in {"why us", "why choose us"}:
        return "why"
    if "closing" in key:
        return "closing"
    compact = key.replace(" ", "")
    if compact in {"services", "ourservices"} or key in {
        "our services",
        "our services and others",
        "services and others",
    }:
        return "services"
    if len(words) <= 3 and words[-1] == "services" and words[0] in {"our", "the"}:
        return "services"
    return None


def looks_like_question(text: str) -> bool:
    stripped = text.strip()
    if stripped.endswith("?"):
        return True
    start = normalize(stripped).split(" ")[:1]
    return start == ["what"] or start == ["how"] or start == ["why"] or start == ["do"] or start == ["can"] or start == ["are"]


def read_paragraphs(path: Path) -> list[tuple[int, str]]:
    document = Document(str(path))
    rows: list[tuple[int, str]] = []
    for paragraph in document.paragraphs:
        text = paragraph_text(paragraph)
        if not text:
            continue
        rows.append((heading_level(paragraph), text))
    if not rows:
        raise ValueError(f"No readable paragraphs in {path}")
    return rows


def next_is_heading(rows: list[tuple[int, str]], index: int) -> bool:
    return index + 1 < len(rows) and rows[index + 1][0] > 0


def collect_bodies(rows: list[tuple[int, str]], start: int) -> tuple[list[str], int]:
    bodies: list[str] = []
    index = start
    while index < len(rows) and rows[index][0] == 0:
        bodies.append(rows[index][1])
        index += 1
    return bodies, index


def parse_outline(path: Path) -> Outline:
    rows = read_paragraphs(path)
    outline = Outline()

    index = 0
    if rows[0][0] == 0:
        # Rare: title stored as a body paragraph.
        outline.title = rows[0][1]
        index = 1
    else:
        outline.title = rows[0][1]
        index = 1
        # Sample Format uses a second Heading 1 as a subtitle; skip it if
        # the following block is still intro copy (body text).
        if index < len(rows) and rows[index][0] == 1 and not classify_section_title(rows[index][1]):
            index += 1

    intro, index = collect_bodies(rows, index)
    outline.intro = intro

    current = "services"
    untitled_blocks = 0

    while index < len(rows):
        level, text = rows[index]
        if level == 0:
            if current == "closing":
                outline.closing_bodies.append(text)
            elif outline.items[current]:
                outline.items[current][-1].bodies.append(text)
            elif current == "services" and not outline.section_titles["services"]:
                outline.intro.append(text)
            index += 1
            continue

        named = classify_section_title(text)
        starts_section = named is not None or next_is_heading(rows, index)

        if current == "faq" and not named and not next_is_heading(rows, index) and not looks_like_question(text):
            current = "closing"
            outline.section_titles["closing"] = text
            index += 1
            bodies, index = collect_bodies(rows, index)
            outline.closing_bodies.extend(bodies)
            continue

        if starts_section and named:
            current = named
            outline.section_titles[current] = text
            index += 1
            continue

        if starts_section and named is None and next_is_heading(rows, index):
            # Unknown section banner (heading immediately followed by another heading).
            untitled_blocks += 1
            current = SECTION_ORDER[min(untitled_blocks - 1, len(SECTION_ORDER) - 1)]
            outline.section_titles[current] = text
            index += 1
            continue

        outline.items[current].append(Item(title=text))
        index += 1
        bodies, index = collect_bodies(rows, index)
        outline.items[current][-1].bodies.extend(bodies)

        if current == "services" and not outline.section_titles["services"]:
            outline.section_titles["services"] = DEFAULT_SECTION_TITLES["services"]

    for key, default in DEFAULT_SECTION_TITLES.items():
        if not outline.section_titles[key] and (outline.items[key] or key == "closing" and outline.closing_bodies):
            outline.section_titles[key] = default

    return outline


def clear_document_body(document: Document) -> None:
    body = document.element.body
    sect_pr = body.find(qn("w:sectPr"))
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def add_paragraph(document: Document, text: str, style: str) -> None:
    paragraph = document.add_paragraph(text)
    try:
        paragraph.style = style
    except KeyError:
        paragraph.style = document.styles[style]


def write_outline(outline: Outline, sample_path: Path, output_path: Path) -> None:
    document = Document(str(sample_path))
    clear_document_body(document)

    add_paragraph(document, outline.title, HEADING_STYLE[1])
    for para in outline.intro:
        add_paragraph(document, para, BODY_STYLE)

    for key in SECTION_ORDER:
        title = outline.section_titles.get(key) or DEFAULT_SECTION_TITLES[key]
        if key != "closing" and not outline.items[key]:
            continue
        if key == "closing" and not outline.closing_bodies and not title:
            continue

        add_paragraph(document, title, HEADING_STYLE[2])
        if key == "closing":
            for para in outline.closing_bodies:
                add_paragraph(document, para, BODY_STYLE)
            continue

        for item in outline.items[key]:
            add_paragraph(document, item.title, HEADING_STYLE[3])
            if item.bodies:
                for para in item.bodies:
                    add_paragraph(document, para, BODY_STYLE)
            else:
                add_paragraph(document, "", BODY_STYLE)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))


def output_name(source: Path, output_dir: Path, in_place: bool) -> Path:
    if in_place:
        return source
    return output_dir / f"{source.stem} - formatted.docx"


def format_file(source: Path, sample: Path, dest: Path) -> Outline:
    outline = parse_outline(source)
    write_outline(outline, sample, dest)
    return outline


def summarize(outline: Outline) -> str:
    parts = [
        f"title={outline.title!r}",
        f"intro={len(outline.intro)}",
    ]
    for key in SECTION_ORDER:
        if key == "closing":
            parts.append(f"closing={len(outline.closing_bodies)}")
        else:
            parts.append(f"{key}={len(outline.items[key])}")
    return ", ".join(parts)


def resolve_sample(path: Path) -> Path | None:
    if path.is_file():
        return path
    folder = path if path.is_dir() else SAMPLE_DIR
    named = folder / DEFAULT_SAMPLE_NAME
    if named.is_file():
        return named
    matches = sorted(
        candidate
        for candidate in folder.glob("*.docx")
        if not candidate.name.startswith("~$")
    )
    return matches[0] if matches else None


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Restyle .docx outlines to match Sample Format.docx heading levels."
    )
    parser.add_argument(
        "--sample",
        type=Path,
        default=SAMPLE_DIR,
        help="Sample outline file, or the sample formats folder (default: sample formats).",
    )
    parser.add_argument(
        "--input-dir",
        type=Path,
        default=INPUT_DIR,
        help="Folder of unformatted .docx files (default: unformatted).",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=OUTPUT_DIR,
        help="Folder for formatted copies (default: formatted).",
    )
    parser.add_argument(
        "--in-place",
        action="store_true",
        help="Overwrite the source files instead of writing copies.",
    )
    parser.add_argument(
        "inputs",
        nargs="*",
        type=Path,
        help="Documents to format. Defaults to all .docx files in the input directory.",
    )
    return parser.parse_args(argv)


def default_inputs(input_dir: Path) -> list[Path]:
    return sorted(
        path
        for path in input_dir.glob("*.docx")
        if not path.name.startswith("~$")
    )


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv if argv is not None else sys.argv[1:])
    sample = resolve_sample(args.sample)
    if sample is None:
        print(f"Sample format not found: {args.sample}", file=sys.stderr)
        return 1

    inputs = list(args.inputs) if args.inputs else default_inputs(args.input_dir)
    if not inputs:
        print(f"No input .docx files found in {args.input_dir}.", file=sys.stderr)
        return 1

    failures = 0
    for source in inputs:
        if source.name.startswith("~$"):
            continue
        if not source.is_file():
            print(f"Skip missing file: {source}", file=sys.stderr)
            failures += 1
            continue
        dest = output_name(source, args.output_dir, args.in_place)
        try:
            outline = format_file(source, sample, dest)
        except Exception as exc:  # noqa: BLE001 — report and continue other files
            print(f"Failed {source}: {exc}", file=sys.stderr)
            failures += 1
            continue
        print(f"Wrote {dest}")
        print(f"  {summarize(outline)}")
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
